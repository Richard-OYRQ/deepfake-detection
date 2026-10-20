#!/usr/bin/env python3
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

def add_h(doc, text, level):
    h = doc.add_heading(text, level=level)
    for run in h.runs: run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)
    return h

def add_p(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.first_line_indent = Cm(0.74)
    run = p.add_run(text)
    run.font.size = Pt(11); run.font.name = '宋体'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    return p

doc = Document()
for s in doc.sections: s.top_margin = Cm(2.54); s.bottom_margin = Cm(2.54); s.left_margin = Cm(3.17); s.right_margin = Cm(3.17)

t = doc.add_heading('DeepFake视频AI鉴伪检测系统\n工作进度报告', level=0)
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
info = doc.add_paragraph(); info.alignment = WD_ALIGN_PARAGRAPH.CENTER
info.add_run('\n报告周期：2026年6月20日—7月5日').font.size = Pt(11)

add_h(doc, '一、工作概述', 1)
add_p(doc, '本阶段围绕DeepFake视频鉴伪系统核心能力建设，完成了系统摸底、提示词迭代、预处理开发、投票策略设计及批量测试全流程工作。系统基于千问qwen3-vl-plus大模型与OpenCV预处理，采用"大模型语义理解+传统CV像素分析"双模态架构，经多轮优化将100个视频测试集检出率从0%提升至86%。')

add_h(doc, '二、主要工作', 1)
add_p(doc, '提示词工程：经历四轮迭代。21项版因"注意力稀释效应"检出率0%，回退至14项提升至56%，18项虚化专项版降至33%，最终14项+预处理版稳定在80%（F1-F10基准集）。验证了"少即是多"的提示词设计原则。')
add_p(doc, '预处理开发：用OpenCV Haar Cascade+拉普拉斯方差量化面部/背景/颈部清晰度，创新引入"面颈比"指标区分换脸蒙版虚化与美颜滤镜。经七轮阈值调优，有效控制了真实视频误报率（0%）并提升换脸检出率。')
add_p(doc, '投票策略：对比1票/2票/3票方案后，设计"1票初筛+3票复检"两阶段策略，在精度与成本间取得平衡（100视频约190次调用）。阶段二成功纠正F6等漏检案例。')
add_p(doc, '模型选型：对比四款千问模型，qwen3-vl-plus综合表现最优，qwen-vl-max过度激进导致高误报。')
add_p(doc, '批量测试：完成4轮对比实验和最终100视频两阶段检测（100/100成功，总耗时63分钟）。')

add_h(doc, '三、关键成果', 1)
add_p(doc, 'F1-F10基准集检出率80%（8/10），100视频批量检出率86%（86/100），真实视频误报率0%。核心创新包括：发现VLM"注意力稀释效应"、确立"量化预处理+精简提示词+多轮投票"组合策略、设计面颈比指标及两阶段检测架构。')

add_h(doc, '四、后续计划', 1)
add_p(doc, '当前主要局限为无法检测零异常的高质量换脸（如F3、F10）。后续拟引入MediaPipe人脸检测增强鲁棒性、增加频域纹理分析检测模式崩溃伪影、探索思维链提示词策略、并建立标准测试集支撑跨模型验证。')

doc.save('D:/DeepFake/DeepFake鉴伪检测工作进度报告_v2.docx')
print('Done')
