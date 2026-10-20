# -*- coding: utf-8 -*-
"""生成《下一步推进工作方案》（约2000字，对照鉴伪实践项目要求）。"""
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn


def set_font(run, name='微软雅黑', size=11, bold=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def h1(doc, text):
    p = doc.add_heading(level=1)
    set_font(p.add_run(text), size=16, bold=True, color=(0x1F, 0x4E, 0x79))


def h2(doc, text):
    p = doc.add_heading(level=2)
    set_font(p.add_run(text), size=13, bold=True, color=(0x2E, 0x74, 0xB5))


def para(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Pt(22)
    set_font(p.add_run(text), size=11)
    return p


def bullet(doc, text):
    p = doc.add_paragraph(style='List Bullet')
    set_font(p.add_run(text), size=11)


def table(doc, headers, rows):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = 'Light Grid Accent 1'
    for j, htext in enumerate(headers):
        cell = t.rows[0].cells[j]
        set_font(cell.paragraphs[0].add_run(htext), size=10, bold=True)
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = t.rows[i + 1].cells[j]
            set_font(cell.paragraphs[0].add_run(str(val)), size=10)
    doc.add_paragraph()


doc = Document()

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_font(p.add_run('AI换脸鉴伪项目：下一步推进工作方案'), size=20, bold=True)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_font(p.add_run('——直面视觉大模型路线困境，对照实践要求的收官设计（2026年8月28日）'), size=12, color=(0x66, 0x66, 0x66))

# ============ 一、现状与困境 ============
h1(doc, '一、现状评估：要求与差距')
para(doc, '对照《鉴伪实践项目要求》，当前状态为：准确率≥85%已达成（v5.4为85.5%、v5.5为97.5%，均在200标注视频全自动口径下）；五类交付物中自动化脚本、运行记录、提示词文档（五版）基本齐备，实验报告与展示PPT需要按最终定稿方案补齐；十个周的进度安排中第一至八周的工作全部完成，第九至十周（提示词深度优化定型+脚本文档收尾）是当前所处位置。')
para(doc, '同时必须直面困境：配对基准实验证明，视觉大模型对高质量1080P换脸零判别力（两个模型各2/6正确），且此前"证据子分"等模型通道信号实为数据源质量混淆的产物。若按"提示词工程为核心方法"的考核口径直接提交以模型票为主判的方案，科学上不诚实；若完全放弃模型通道，又偏离实践要求的主题。本方案的设计目标是在两者之间找到既诚实又可交付的平衡：模型通道保持"结构化证据提供者"角色，判定主信号转向配对实验验证过的CV内容特征，提示词工程的全部探索历程作为核心交付叙事。')

# ============ 二、总体思路 ============
h1(doc, '二、总体思路：三线并进')
para(doc, '第一条线（科学主线）：在14对配对基准上验证"面部软化系"特征的方向一致性——配对实验中已发现面部清晰度、面背清晰比、面颈比、噪声残留比四个特征在3对样本上方向100%一致（假脸更低），这是当前唯一在干净数据上表现出换脸自身指纹的信号。若14对验证通过，将其纳入判定策略作为"换脸软化指纹"证据。')
para(doc, '第二条线（交付主线）：把v5.4/v5.5系统按实践要求定稿。提示词文档以"16项主提示词+14项控方提示词+证据图注入"三件套为最终版，逐版记录调优思路；运行记录整理历轮检测对比表；实验报告如实呈现"模型通道裁决实验"与"配对基准证伪"的完整证据链——提示词工程的真实价值恰恰体现在它精确测量出了模型边界。')
para(doc, '第三条线（探索主线）：为视觉外观通道之外的新信号做可行性验证。音视频唇同步特征（PyAV已具备解码条件）的第一版实验已证伪，但升级版（嘴部暗腔面积时序等更精细的口型信号）尚未尝试；频域指纹与压缩痕迹分析尚未开发。这两类信号若有效，可与"面部软化系"共同构成不依赖源质量的内容特征组。')

# ============ 三、具体任务 ============
h1(doc, '三、具体任务与时间表')
h2(doc, '任务一：14对全基准"面部软化系"验证（第1天，纯本地零成本）')
bullet(doc, '对14对配对样本提取全部CV特征，统计四个软化系特征的方向一致率；')
bullet(doc, '若一致率≥80%：设计"软化指纹"判定规则（如面背比降幅阈值），在14对上测单特征准确率；')
bullet(doc, '产出：配对特征对照表与初步判定规则草案。')
h2(doc, '任务二：最终版提示词定稿与配对定型实验（第2-3天，约1.5元API成本）')
bullet(doc, '在14对配对基准上跑提示词变体对照：现行16项版、"软化指纹数据注入版"（把CV软化特征作为量化证据写入提示词附录）、"局部放大指引版"；')
bullet(doc, '以配对准确率与逐项证据分离度为准绳，选定最终版提示词并定稿；')
bullet(doc, '产出：最终版提示词文档（含每版调优思路与调整细节，满足要求中的提示词文档交付物）。')
h2(doc, '任务三：200标注集终裁与全量回归（第3天，约25元API成本）')
bullet(doc, '将定稿策略（软化指纹+现有v5.4路由+三分类）应用到200标注集，全量实跑一次作为官方留档数据；')
bullet(doc, '产出：最终混淆矩阵、误报/漏检清单、优化前后准确率对比表（满足运行记录交付物）。')
h2(doc, '任务四：报告与PPT产出（第4-5天）')
bullet(doc, '实验报告：API调用过程、提示词五版演进、三轮验证与配对证伪、最终方案总结；')
bullet(doc, '展示PPT：系统架构图、性能演进图（已备matplotlib图表）、配对实验的"信号证伪"叙事、心得体会；')
bullet(doc, '脚本说明文档：main.py、build_feature_table.py、analyze_features.py、各验证脚本的使用说明（满足自动化脚本交付物）。')

# ============ 四、风险与备选 ============
h1(doc, '四、风险与备选方案')
bullet(doc, '风险一：软化系特征在14对全基准上方向不一致。备选：判定策略退回v5.4（85.5%）作为最终交付版本，报告如实说明模型通道边界，将软化指纹列为展望；')
bullet(doc, '风险二：考核口径要求"以模型判定为主"。备选：在形式上保持模型为最终判定者（16项提示词+全部CV证据注入后由模型出结论），判定规则作为提示词内的量化指导——这是当前架构已经满足的形式；')
bullet(doc, '风险三：qwen-vl-max退役（2026-10-10）影响演示。备选：最终方案锁定qwen3-vl-plus（现役），退役模型仅作对照实验记录；')
bullet(doc, '风险四：全量实跑费用与时间。备选：以离线重算（reapply_fusion+融合先验表）作为官方数据，实跑仅作演示录制。')

# ============ 五、预期产出 ============
h1(doc, '五、预期产出清单（对照要求逐项）')
table(doc,
      ['要求交付物', '状态', '补齐动作'],
      [
          ['实验报告', '部分（多份阶段报告）', '任务四整合为最终版'],
          ['展示PPT', '未产出', '任务四新制'],
          ['自动化脚本', '已具备（主程序+特征表+分析工具）', '任务四补说明文档'],
          ['提示词文档（初始/优化/最终）', '五版提示词已存档', '任务二定稿最终版'],
          ['运行记录（对比表/清单）', '已具备（历轮results+报告）', '任务三补最终轮'],
          ['准确率≥85%', '已达成（85.5%/97.5%）', '任务三实跑留档'],
      ])

doc.save(r'D:\deepfake\AI换脸鉴伪项目下一步工作方案_20260828.docx')
print('saved')
