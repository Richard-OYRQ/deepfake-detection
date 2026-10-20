# -*- coding: utf-8 -*-
"""汇总生成《视频深度伪造鉴伪课题综合报告》docx（约30000字）。"""
import sys

sys.path.insert(0, r"E:\deepfake")
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from report_content_part1 import SECTIONS as P1
from report_content_part2 import SECTIONS as P2
from report_content_part3 import SECTIONS as P3
from report_content_part4 import SECTIONS as P4

OUT = r"E:\deepfake\研究文档\20260929_课题综合报告.docx"

doc = Document()
for sec in doc.sections:
    sec.top_margin = Cm(2.5); sec.bottom_margin = Cm(2.5)
    sec.left_margin = Cm(2.8); sec.right_margin = Cm(2.8)


def set_font(run, name="宋体", size=12, bold=False, color=None):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts"); rPr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), name)


def title(text):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run(text), name="黑体", size=16, bold=True)


def subtitle(text):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run(text), size=10.5)


def h1(text):
    p = doc.add_paragraph()
    set_font(p.add_run(text), name="黑体", size=14, bold=True, color=(0x1F, 0x4E, 0x79))
    p.paragraph_format.space_before = Pt(10); p.paragraph_format.space_after = Pt(4)


def h2(text):
    p = doc.add_paragraph()
    set_font(p.add_run(text), name="黑体", size=12, bold=True, color=(0x2E, 0x74, 0xB5))
    p.paragraph_format.space_before = Pt(6); p.paragraph_format.space_after = Pt(3)


def para(text, indent=True):
    p = doc.add_paragraph()
    if indent:
        p.paragraph_format.first_line_indent = Pt(24)
    p.paragraph_format.line_spacing = 1.4
    set_font(p.add_run(text))


title("视频深度伪造鉴伪课题综合报告")
subtitle("——系统设计、提示词工程、实验方法与全球视觉大模型能力对比（2026年9月29日）")

total = 0
for sections in (P1, P2, P3, P4):
    for level, head, paras in sections:
        (h1 if level == 1 else h2)(head)
        total += len(head)
        for t in paras:
            para(t, indent=(head != "参考文献"))
            total += len(t)

doc.save(OUT)
print(f"已生成: {OUT}")
print(f"正文总字数（含标题）: {total}")
