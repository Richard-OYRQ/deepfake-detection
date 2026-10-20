# -*- coding: utf-8 -*-
"""生成《外部评测集分层抽样实验报告》Word 文档。"""
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()
for sec in doc.sections:
    sec.top_margin = Cm(2.5); sec.bottom_margin = Cm(2.5)
    sec.left_margin = Cm(2.8); sec.right_margin = Cm(2.8)


def set_font(run, name="宋体", size=12, bold=False, color=None):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), name)


def title(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run(text), name="黑体", size=16, bold=True)
    p.paragraph_format.space_after = Pt(6)


def h1(text):
    p = doc.add_paragraph()
    set_font(p.add_run(text), name="黑体", size=14, bold=True, color=(0x1F, 0x4E, 0x79))
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)


def h2(text):
    p = doc.add_paragraph()
    set_font(p.add_run(text), name="黑体", size=12, bold=True, color=(0x2E, 0x74, 0xB5))
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(3)


def para(text):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Pt(24)
    p.paragraph_format.line_spacing = 1.4
    set_font(p.add_run(text))
    return p


def table(headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    for i, htxt in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = ""
        set_font(c.paragraphs[0].add_run(htxt), name="黑体", size=10, bold=True)
    for row in rows:
        cells = t.add_row().cells
        for i, ctxt in enumerate(row):
            cells[i].text = ""
            set_font(cells[i].paragraphs[0].add_run(str(ctxt)), size=9.5)
    doc.add_paragraph()


title("外部评测集分层抽样实验报告")
sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_font(sub.add_run("实验日期：2026-09-13 至 09-14 ｜ 检测系统：v5.7（qwen3-vl-plus 三件套）｜ 报告日期：2026-09-14"),
         size=9, color=(0x80, 0x80, 0x80))
doc.add_paragraph()

# 一
h1("一、实验背景与目的")
para("此前系统性能基准（自建 200 视频：全自动综合正确率 97.5%、硬误报 0、硬漏检 1）是在自建数据集上"
     "校准得出的，该数据集伪造样本全部不高于 720P。为检验系统在分布外样本上的泛化能力，"
     "项目构建了 1000 个外部评测视频（500 真实 + 500 伪造，覆盖 480P/720P/1080P 与 7 类伪造流派），"
     "并开展三组策略优化实验与一次分层抽样评测，回答两个问题：一是系统对高清伪造的检测能力如何；"
     "二是候选优化策略能否改善高清伪造检出而不牺牲真实误报。")

# 二
h1("二、外部评测集构建")
para("评测集共 1000 个视频，来源两部分：angads24/deepfake-video（200 真实 + 200 伪造，全 1080P，"
     "FF++ 衍生）与 bitmind/FaceForensicsC23（300 真实 + 300 伪造，6 种伪造方法各 50，分辨率 480P-1080P 混合）。"
     "全部视频经 ffprobe 探测分辨率与全帧解码校验，无损坏文件。")
table(["构成", "数量", "分辨率特征"], [
    ["真实视频", "500", "1080P 238 / 720P 88 / 480P 系 170 / 其他 4"],
    ["伪造-DeepFakeDetection", "50", "全 1080P"],
    ["伪造-Deepfakes", "50", "1080P 7 / 720P 16 / 480P 系 27"],
    ["伪造-Face2Face", "50", "1080P 6 / 720P 17 / 480P 系 27"],
    ["伪造-FaceShifter", "50", "1080P 6 / 720P 10 / 480P 系 34"],
    ["伪造-FaceSwap", "50", "1080P 5 / 720P 18 / 480P 系 27"],
    ["伪造-NeuralTextures", "50", "1080P 2 / 720P 19 / 480P 系 29"],
    ["伪造-未标注（FF++衍生）", "200", "全 1080P"],
])

# 三
h1("三、基线实验：14 个 1080P 配对样本")
para("从评测集抽取 7 个 1080P 真实与 7 个 1080P Deepfakes 伪造组成配对样本，全部进入高清"
     "模型主导分支（保守票直接终裁、跳过控方轮、CV 冲突豁免）。基线结果：真实组 7 个全部判定"
     "无伪造（误报 0），伪造组 6 个漏检 + 1 个存疑（检出 0）。")
h2("3.1 判定结果")
table(["组别", "模型票", "置信度", "异常项", "CV先验", "终裁"], [
    ["真实 ×7", "全'无伪造'", "94-96", "全 0/16", "0.29-0.66", "无伪造 7（误报 0）"],
    ["伪造 ×6", "全'无伪造'", "94-96", "全 0/16", "0.30-0.55", "无伪造 6（漏检）"],
    ["伪造 894_848", "'换脸'", "96", "#15/#16 双异常", "0.546", "存疑（先验背书 0.70 拦截）"],
])
h2("3.2 逐项判定矩阵")
para("真假两组的逐项判定模式几乎完全一致：#1 边缘模糊、#3 皮肤平滑、#4 肤色差异、#7 毛发边缘、"
     "#11 帧间波动、#12 面颈差异六项软证据在 14 个视频上全覆盖触发，硬证据项全部零触发"
     "（除 894_848 的 #15/#16）。生物信号三项（#6/#9/#10）两组全部正常。")
h2("3.3 关键发现")
para("其一，模型在 1080P 上对 Deepfakes 流派伪影视觉失明，六项软证据的触发被系统性归因于美颜滤镜，"
     "软证据豁免机制成为漏检通道。其二，置信度在真假组间无判别力（同为 94-96）。其三，CV 先验"
     "（以低清伪造校准）对高清伪造失效，两组先验分布重叠且伪造组整体略低。其四，894_848 在"
     "重跑中模型票从'换脸'翻转为'无伪造'，暴露高清分支单票直接终裁的随机脆弱性。")

# 四
h1("四、三组优化实验")
h2("4.1 实验 1：先验背书阈值 0.70→0.50")
para("动机：894_848 的换脸票（96 置信）被 0.70 背书拦截，降档可使其定罪。结果：84.mp4（历史误报"
     "样本）保持存疑无回归；但 894_848 模型票随机翻转为'无伪造'，实验未得到预期验证。该改动在"
     "后续 100 视频评测中被证明造成 5 个真实误报（见第五章），已回滚。")
h2("4.2 实验 2：高清分支恢复 CV 可视化核验")
para("动机：热力图与面颈剖面曲线在低清分支有效（配对实验敏感 45%/特异 93%），高清分支此前豁免。"
     "结果：14 个视频核验输出 13 个'反对'、1 个未输出，检出与误报均无变化——模型在高清样本上"
     "未确认任何面颈纹理跳变。改动已回滚。")
h2("4.3 实验 3：提示词新增 GAN 重建指纹维度（17 项版）")
para("动机：16 项维度无 Deepfakes 流派专属检测项，新增第 17 项'五官重建痕迹与纹理噪声模式'，"
     "并配套美颜不可解释清单、软证据归因义务、生物信号例外条款。结果：#17 在 14 个视频上全部判"
     "'正常'，零触发，检出与误报均无变化——文字提示无法让模型看见它感知不到的重建伪影。")

# 五
h1("五、分层抽样 100 视频评测（合并策略）")
para("按类别×分辨率分层抽取 100 个：真实 50（高清 25 + 低清 25）、伪造 50（6 方法各 7 + 未标注 8），"
     "以三实验合并策略（backstop 0.50 + 高清 CV 核验 + 17 项提示词）全量检测，10 批短进程完成，"
     "100/100 运行成功。")
h2("5.1 总体指标")
table(["组别", "检出/误报", "存疑", "漏检/正确"], [
    ["伪造 50", "检出 2（4%）", "15（30%）", "漏检 33（66%）"],
    ["真实 50", "误报 5（10%）", "14（28%）", "正确 31（62%）"],
])
h2("5.2 分层明细")
table(["分层", "结果", "评价"], [
    ["高清伪造 30", "检出 0，存疑 1，漏 29", "黑洞依旧，三实验全部无效"],
    ["低清伪造 20", "检出 2，存疑 14，漏 4", "可接受（80% 不硬漏）"],
    ["高清真实 25", "误报 0，存疑 0", "完美"],
    ["低清真实 25", "误报 5，存疑 14，正确 6", "严重退化（误报 20%）"],
])
h2("5.3 误报归因")
para("5 个误报全部为低清真实视频，其 CV 先验分别为 0.503、0.520、0.562、0.667、0.676——"
     "全部落在 0.50-0.70 区间。backstop 从 0.70 降至 0.50 恰好放行这批样本的定罪路径。"
     "在基线 0.70 阈值下，它们均会被拦截为存疑，误报为 0。据此认定实验 1 的降档是误报增长"
     "的直接原因，予以回滚。")

# 六
h1("六、结论与决策")
para("第一，高清 Deepfakes 伪造是当前系统在提示词与判据层面的无解盲区：模型视觉失明、CV 先验失效、"
     "证据图不触发，三组优化实验（阈值降档、通道开启、提示词扩展）均无法改善 30 个高清伪造的 0 检出。"
     "这印证了提示词工程的能力边界——提示词只能组织模型已有的能力，不能创造它没有的感知。")
para("第二，合并策略总体判定为不建议采用：检出收益仅 2 个（全部来自低清），误报代价 5 个（低清真实）。"
     "已将 backstop 回滚至 0.70、CV 核验与提示词改动全部回滚，系统恢复 v5.7 基线。")
para("第三，后续改进方向按优先级：一是 CV 侧修复——将高清伪造样本纳入先验模型校准，验证先验对"
     "高清伪造是否系统性偏低，必要时构建高清伪造专用先验；二是课题三 CNN 像素级特征检测——"
     "本评测集的 30 个高清伪造漏检样本与 5 个低清真实误报样本可直接作为训练难例与验证集；"
     "三是高清分支的单票随机性问题——探索保守票终裁前的最小复检机制（以 894_848 票翻转的"
     "两次实测为证据）。")

# 七
h1("七、附录：实验配置与复现")
para("基线配置：config.yaml（v5.7，backstop 0.70）。实验 1：config_exp1.yaml；实验 2：processor.py "
     "两处高清豁免临时移除；实验 3：config_exp3.yaml（17 项提示词，见 prompt_v2_exp3.py）；"
     "分层评测：config_exp4c1-c10.yaml（10 批×10 视频）。结果目录：output_bench7、output_bench7t、"
     "output_exp1、output_exp2、output_exp3、output_exp4c1-c10。数据：外部评测集100/"
     "（1000 视频，分类清单见《外部评测集1000分类清单与流派目录》）。")

doc.save(r"E:\DeepFake\外部评测实验报告.docx")
print("DOCX saved")
