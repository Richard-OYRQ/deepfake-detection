# -*- coding: utf-8 -*-
"""生成证据图读图指南（表格版，无装饰符号，可直接放进文档）。"""
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = r"E:\deepfake\研究文档\证据图读图指南_表格版.docx"

doc = Document()
for sec in doc.sections:
    sec.top_margin = Cm(2.5); sec.bottom_margin = Cm(2.5)
    sec.left_margin = Cm(2.2); sec.right_margin = Cm(2.2)

def set_font(run, name="宋体", size=11, bold=False, color=None):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts"); rPr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), name)

p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_font(p.add_run("证据图读图指南"), name="黑体", size=15, bold=True)

headers = ["证据图", "图里有什么", "正常特征", "可疑特征（换脸痕迹）"]
rows = [
    ["纹理密度热力图", "绿框为人脸、红框为颈部；暖色表示纹理细节多，冷色表示平滑", "脸与脖子颜色接近，过渡自然", "脸冷脖热：面部大片冷色、颈部暖色，说明面部被重绘平滑化"],
    ["面颈剖面图", "曲线纵轴为纹理强度，绿线为面颈分界，左侧面部、右侧颈部", "曲线平缓连续，面颈强度接近", "绿线处曲线陡降，面部纹理强度明显低于颈部"],
    ["遮挡交界放大图", "字幕、话筒、手等物体穿过人脸的画面放大", "遮挡物边缘完整、锐利、连续", "交界处涂抹、断裂、笔画变形、边缘融化"],
    ["转动序列对比图", "转头前2帧与转回后2帧并排拼接", "五官相对位置连续，只有角度变化", "五官跳动、脸型抖动、身份闪回"],
    ["帧差增强图", "前后两帧相减后放大，高亮为变化大的区域", "高亮集中在嘴、眼等正常运动区", "高亮杂乱洒满面部区域"],
]
t = doc.add_table(rows=1 + len(rows), cols=4)
t.style = "Table Grid"
for j, h in enumerate(headers):
    c = t.rows[0].cells[j]; c.text = ""
    set_font(c.paragraphs[0].add_run(h), name="黑体", size=11, bold=True)
for i, row in enumerate(rows, 1):
    for j, v in enumerate(row):
        c = t.rows[i].cells[j]; c.text = ""
        set_font(c.paragraphs[0].add_run(str(v)), size=11)

p = doc.add_paragraph()
set_font(p.add_run("读图要点", ), name="黑体", size=12, bold=True)
for text in [
    "1. 先看时间戳（绿字t等于具体秒数），回到原视频同一时刻对照查看。",
    "2. 选帧本身是机器判断，程序挑出的是最可疑的瞬间，人只需要裁决。",
    "3. 建议对比真实视频与伪造视频的同名证据图，感受正常与可疑的差异。",
    "4. 任何一张图的异常都不单独定罪，多张图互相印证才构成证据链。",
]:
    p = doc.add_paragraph()
    set_font(p.add_run(text), size=11)

doc.save(OUT)
print("已生成:", OUT)
